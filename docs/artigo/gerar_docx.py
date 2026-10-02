#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Monta o artigo do TCC em .docx a partir dos arquivos Markdown de cada seção.

Usa o TEMPLATE_ARTIGO_TCC_SPTECH.docx como doador de formatação: abre o
template, esvazia o corpo (preservando cabeçalho, rodapé e configuração de
página) e reescreve o conteúdo aplicando a formatação direta que o template usa.

Por que formatação direta e não estilos: o template **não define estilos de
título** — "1 Introdução" e um parágrafo comum são ambos `Normal`, diferindo
apenas por negrito, caixa alta, alinhamento e `outlineLvl`. Aplicar `Heading 1`
produziria um documento azul e com fonte diferente do modelo da faculdade.

Uso, a partir da raiz do repositório:
    ml-service/venv/Scripts/python.exe docs/artigo/gerar_docx.py

Dependência: python-docx (ver requirements-artigo.txt). Não faz parte do
runtime de nenhum serviço.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

RAIZ = Path(__file__).resolve().parents[2]          # .../tcc-stocksense/code
ARTIGO = RAIZ / "docs" / "artigo"
TEMPLATE = RAIZ.parent.parent / "docs" / "exemplos" / "TEMPLATE_ARTIGO_TCC_SPTECH.docx"
SAIDA = ARTIGO / "artigo-stocksense.docx"

# Ordem de montagem. Arquivos ausentes são pulados com aviso.
SECOES = [
    "00-frontmatter.md",
    "01-introducao.md",
    "02-fundamentacao.md",
    "03-metodologia.md",
    "04-resultados.md",
    "05-consideracoes.md",
    "referencias.md",
]

LARGURA_FIGURA = Cm(15.0)   # largura útil da página (A4 menos margens de 3 cm)


# ─────────────────────────────────────────────────────────────────────────────
# Formatação — espelha exatamente o que o template aplica em cada tipo de bloco
# ─────────────────────────────────────────────────────────────────────────────
def _caps(run) -> None:
    """Liga <w:caps/>, que o template usa nos títulos de seção."""
    run._element.get_or_add_rPr().append(run._element.makeelement(qn("w:caps"), {}))


def _nivel(par, nivel: int) -> None:
    """Grava o outlineLvl, que alimenta o painel de navegação e o sumário."""
    el = par._p.get_or_add_pPr().makeelement(qn("w:outlineLvl"), {qn("w:val"): str(nivel)})
    par._p.get_or_add_pPr().append(el)


def _base(par, *, centro=False, keep=False) -> None:
    """Entrelinha 1,5, recuo direito -1 twip e justificação — o padrão do corpo."""
    pf = par.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(0)
    pf.right_indent = Pt(-0.05)
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if centro else WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.keep_with_next = keep


def titulo(doc, texto: str, nivel: int):
    """Título de seção: negrito + caixa alta, justificado, mantido com o próximo."""
    p = doc.add_paragraph()
    _base(p, keep=True)
    _nivel(p, nivel - 1)
    r = p.add_run(texto)
    r.bold = True
    _caps(r)
    return p


def bloco_letra(doc, texto: str):
    """Bloco 'A. Componente' da seção de arquitetura: negrito, sem caixa alta."""
    p = doc.add_paragraph()
    _base(p, keep=True)
    _nivel(p, 3)
    p.add_run(texto).bold = True
    return p


def legenda(doc, texto: str, keep=False):
    """Legenda de figura/tabela e linha de fonte: centralizada e em itálico."""
    p = doc.add_paragraph()
    _base(p, centro=True, keep=keep)
    r = p.add_run(texto)
    r.italic = True
    return p


def equacao(doc, texto: str):
    """Equação numerada: centralizada, sem ênfase."""
    p = doc.add_paragraph()
    _base(p, centro=True)
    p.add_run(texto)
    return p


_INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")


def corpo(doc, texto: str, marcador: str | None = None):
    """
    Parágrafo comum, convertendo **negrito**, *itálico* e `código` em runs.

    `marcador` transforma o parágrafo em item de lista. O template não define
    os estilos `List Number`/`List Bullet` do Word, então o marcador é escrito
    como texto e o recuo é aplicado à mão — o que também evita que o Word
    renumere de forma inesperada entre seções.
    """
    p = doc.add_paragraph()
    _base(p)
    if marcador:
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.add_run(marcador + " ")
    for pedaco in _INLINE.split(texto):
        if not pedaco:
            continue
        if pedaco.startswith("**") and pedaco.endswith("**"):
            p.add_run(pedaco[2:-2]).bold = True
        elif pedaco.startswith("*") and pedaco.endswith("*"):
            p.add_run(pedaco[1:-1]).italic = True
        elif pedaco.startswith("`") and pedaco.endswith("`"):
            r = p.add_run(pedaco[1:-1])
            r.font.name = "Courier New"
            r.font.size = Pt(10)
        else:
            p.add_run(pedaco)
    return p


def _bordas(t) -> None:
    """Aplica bordas finas à tabela — o template só traz a 'Normal Table', sem grade."""
    props = t._tbl.tblPr
    borders = props.makeelement(qn("w:tblBorders"), {})
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.makeelement(qn(f"w:{lado}"), {
            qn("w:val"): "single", qn("w:sz"): "4", qn("w:color"): "000000"})
        borders.append(el)
    props.append(borders)


def tabela(doc, linhas: list[str]):
    """Converte um bloco Markdown de tabela em tabela do Word."""
    celulas = [[c.strip() for c in l.strip().strip("|").split("|")] for l in linhas]
    cabecalho, dados = celulas[0], celulas[2:]        # [1] é a linha de alinhamento
    t = doc.add_table(rows=1, cols=len(cabecalho))
    _bordas(t)
    for i, texto in enumerate(cabecalho):
        cel = t.rows[0].cells[i]
        cel.text = ""
        p = cel.paragraphs[0]
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(texto.replace("**", ""))
        r.bold = True
        r.font.size = Pt(10)
    for linha in dados:
        cels = t.add_row().cells
        for i, texto in enumerate(linha[:len(cabecalho)]):
            cels[i].text = ""
            p = cels[i].paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(texto.replace("**", ""))
            r.font.size = Pt(10)
    return t


def figura(doc, spec: str):
    """
    Insere a imagem referenciada por [FIGURA: ...] ou, se ela não existir,
    um marcador visível para que a lacuna não passe despercebida na revisão.
    """
    caminho = spec.split("—")[0].strip()
    for base in (RAIZ, RAIZ.parent, RAIZ.parent.parent):
        alvo = base / caminho
        if alvo.is_file() and alvo.suffix.lower() in (".png", ".jpg", ".jpeg"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(alvo), width=LARGURA_FIGURA)
            return True
    p = doc.add_paragraph()
    _base(p, centro=True)
    r = p.add_run(f"[FIGURA PENDENTE: {caminho}]")
    r.bold = True
    return False


# ─────────────────────────────────────────────────────────────────────────────
# Parser do Markdown — reconhece só os blocos que os arquivos do artigo usam
# ─────────────────────────────────────────────────────────────────────────────
RE_FIG = re.compile(r"^`\[FIGURA:\s*(.+?)\]`$", re.S)
RE_FIG_ABRE = re.compile(r"^`\[FIGURA:")
RE_LEGENDA = re.compile(r"^\*\*(Figura|Tabela) \d+ - .+\*\*$")
RE_LETRA = re.compile(r"^\*\*[A-Z]\. .+\*\*$")
RE_EQ = re.compile(r"^\s{4,}\S.*\(\d+\)\s*$")
RE_LISTA = re.compile(r"^(\d+\.|[-*])\s+(.*)$")


def montar_secao(doc, caminho: Path, relatorio: list[str]) -> None:
    linhas = caminho.read_text(encoding="utf-8").splitlines()
    i, n = 0, len(linhas)
    figuras = pendentes = tabelas = 0

    while i < n:
        linha = linhas[i]
        cru = linha.rstrip()
        texto = cru.strip()

        if not texto:
            i += 1
            continue

        if texto.startswith("<!--"):                       # comentário de autoria
            while i < n and "-->" not in linhas[i]:
                i += 1
            i += 1
            continue

        if texto.startswith("#"):                          # título
            nivel = len(texto) - len(texto.lstrip("#"))
            titulo(doc, texto.lstrip("#").strip(), nivel)
            i += 1
            continue

        if RE_LEGENDA.match(texto):                        # legenda de figura/tabela
            legenda(doc, texto.strip("*"), keep=True)
            i += 1
            continue

        if texto.startswith("Fonte:"):
            legenda(doc, texto)
            i += 1
            continue

        if RE_FIG_ABRE.match(texto):
            # O marcador pode ocupar mais de uma linha quando o caminho é longo.
            spec = [texto]
            while i < n and "]`" not in linhas[i]:
                i += 1
                if i < n:
                    spec.append(linhas[i].strip())
            i += 1
            m = RE_FIG.match(" ".join(spec))
            if m and figura(doc, m.group(1)):
                figuras += 1
            else:
                pendentes += 1
            continue

        if RE_LETRA.match(texto):                          # bloco "A. Componente"
            bloco_letra(doc, texto.strip("*"))
            i += 1
            continue

        if RE_EQ.match(cru):                               # equação numerada
            equacao(doc, texto)
            i += 1
            continue

        if texto.startswith("|"):                          # tabela
            bloco = []
            while i < n and linhas[i].strip().startswith("|"):
                bloco.append(linhas[i])
                i += 1
            if len(bloco) >= 2:
                tabela(doc, bloco)
                tabelas += 1
            continue

        m = RE_LISTA.match(texto)
        if m:                                              # item de lista
            conteudo = [m.group(2)]
            i += 1
            while i < n and linhas[i].startswith("   ") and linhas[i].strip() \
                    and not RE_EQ.match(linhas[i].rstrip()):
                conteudo.append(linhas[i].strip())
                i += 1
            marcador = m.group(1) if m.group(1)[0].isdigit() else "•"
            corpo(doc, " ".join(conteudo), marcador=marcador)
            continue

        if texto.startswith(">"):                          # nota de trabalho: não entra
            i += 1
            continue

        paragrafo = [texto]                                # parágrafo comum
        i += 1
        while i < n:
            prox = linhas[i].strip()
            if (not prox or prox.startswith(("#", "|", ">", "`"))
                    or RE_LEGENDA.match(prox) or RE_LETRA.match(prox)
                    or RE_LISTA.match(prox) or prox.startswith("Fonte:")
                    or RE_EQ.match(linhas[i].rstrip())):
                break
            paragrafo.append(prox)
            i += 1
        corpo(doc, " ".join(paragrafo))

    relatorio.append(
        f"  {caminho.name}: {figuras} figuras, {pendentes} pendentes, {tabelas} tabelas")


def main() -> int:
    if not TEMPLATE.is_file():
        print(f"ERRO: template não encontrado em {TEMPLATE}")
        return 1

    doc = Document(str(TEMPLATE))

    # Esvazia o corpo preservando o sectPr (margens, cabeçalho e rodapé vivem nele).
    corpo_xml = doc.element.body
    for filho in list(corpo_xml):
        if not filho.tag.endswith("}sectPr"):
            corpo_xml.remove(filho)

    relatorio, ausentes = [], []
    for nome in SECOES:
        caminho = ARTIGO / nome
        if caminho.is_file():
            montar_secao(doc, caminho, relatorio)
        else:
            ausentes.append(nome)

    doc.save(str(SAIDA))

    print(f"Gerado: {SAIDA.relative_to(RAIZ)}")
    for l in relatorio:
        print(l)
    if ausentes:
        print("  seções ainda não escritas: " + ", ".join(ausentes))
    print("\nConfira no Word: títulos em caixa alta, entrelinha 1,5, "
          "legendas centralizadas em itálico e o cabeçalho com o DOI.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
